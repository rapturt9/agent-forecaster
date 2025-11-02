"""Word document generation for forecasting reports."""

from typing import Any
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ForecastDocumentGenerator:
    """Generates Word documents for forecasting reports."""

    def __init__(self):
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """Set up document styles."""
        # Set default font
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

    def add_title(self, title: str):
        """Add document title."""
        heading = self.doc.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_metadata(self, question: str, question_type: str, timestamp: str, context: str = None):
        """Add forecast metadata."""
        self.doc.add_paragraph(f"Question: {question}", style='Heading 2')
        self.doc.add_paragraph(f"Type: {question_type}")
        self.doc.add_paragraph(f"Generated: {timestamp}")
        if context:
            self.doc.add_paragraph(f"Context: {context}")
        self.doc.add_paragraph()  # Blank line

    def add_section(self, title: str, content: str, level: int = 1):
        """Add a section with heading and content."""
        self.doc.add_heading(title, level=level)
        self.doc.add_paragraph(content)
        self.doc.add_paragraph()  # Blank line

    def add_forecast_result(self, forecast: str, confidence: str, reasoning: str):
        """Add forecast result section."""
        self.doc.add_heading("Forecast Result", level=1)

        # Forecast value
        p = self.doc.add_paragraph()
        p.add_run("Forecast: ").bold = True
        p.add_run(str(forecast))

        # Confidence
        p = self.doc.add_paragraph()
        p.add_run("Confidence: ").bold = True
        p.add_run(str(confidence))

        # Reasoning
        self.doc.add_heading("Reasoning", level=2)
        self.doc.add_paragraph(reasoning)
        self.doc.add_paragraph()

    def add_model(self, model_id: int, name: str, code: str, description: str = ""):
        """Add a forecasting model."""
        self.doc.add_heading(f"Model {model_id}: {name}", level=2)

        if description:
            self.doc.add_paragraph(description)

        # Add Squiggle code block
        self.doc.add_paragraph("Squiggle Code:", style='Heading 3')
        code_para = self.doc.add_paragraph(code)
        code_para.style = 'Normal'
        # Make code monospace
        for run in code_para.runs:
            run.font.name = 'Courier New'
            run.font.size = Pt(9)

        self.doc.add_paragraph()  # Blank line

    def add_models_section(self, models: list):
        """Add all forecasting models."""
        self.doc.add_heading("Forecasting Models", level=1)

        for model in models:
            self.add_model(
                model_id=model.get('model_id', 0),
                name=model.get('name', 'Unknown Model'),
                code=model.get('code', ''),
                description=model.get('description', '')
            )

    def add_research_summary(self, tool_calls: int, duration: float, research_steps: list):
        """Add research process summary."""
        self.doc.add_heading("Research Process", level=1)

        # Summary stats
        p = self.doc.add_paragraph()
        p.add_run("Total Tool Calls: ").bold = True
        p.add_run(str(tool_calls))

        p = self.doc.add_paragraph()
        p.add_run("Duration: ").bold = True
        p.add_run(f"{duration:.1f} seconds")

        # Research steps (simplified)
        if research_steps:
            self.doc.add_heading("Key Research Steps", level=2)
            for step in research_steps[:10]:  # Limit to first 10
                if step.get('step_type') == 'tool_call':
                    self.doc.add_paragraph(
                        f"• {step.get('content', '')}",
                        style='List Bullet'
                    )

    def save(self, filepath: str):
        """Save document to file."""
        Path(filepath).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(filepath)
        return filepath


def generate_forecast_document(
    output_path: str,
    question: str,
    question_type: str,
    forecast: str,
    confidence: str,
    reasoning: str,
    models: list = None,
    context: str = None,
    tool_calls: int = 0,
    duration: float = 0,
    research_steps: list = None
) -> str:
    """
    Generate a complete forecast Word document.

    Args:
        output_path: Path to save the document
        question: The forecasting question
        question_type: Type of question (binary, numerical, categorical)
        forecast: The forecast value
        confidence: Confidence level
        reasoning: Detailed reasoning
        models: List of forecasting models with Squiggle code
        context: Optional context for the forecast
        tool_calls: Number of tool calls made
        duration: Time taken for forecast
        research_steps: List of research steps taken

    Returns:
        Path to generated document
    """
    gen = ForecastDocumentGenerator()

    # Add title
    gen.add_title("AI Forecasting Report")

    # Add metadata
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    gen.add_metadata(question, question_type, timestamp, context)

    # Add forecast result
    gen.add_forecast_result(forecast, confidence, reasoning)

    # Add models if provided
    if models:
        gen.add_models_section(models)

    # Add research summary
    if research_steps:
        gen.add_research_summary(tool_calls, duration, research_steps)

    # Save
    return gen.save(output_path)
